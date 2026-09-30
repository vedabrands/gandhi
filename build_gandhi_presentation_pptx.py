"""
build_gandhi_presentation_pptx.py
Converts the Gandhi Jayanti 10-slide medieval manuscript presentation into
a standalone PowerPoint (.pptx) presentation with full design fidelity:
- 16:9 Widescreen (13.333" x 7.5")
- Exact Aged Parchment (#FAF4E6 / #F1E4C3), Leather Brown (#2A1A0E), Deep Maroon (#7A1F1F), Antique Gold (#B8862B) theme
- Double-border parchment sheet, ornate gold corner flourishes (❖)
- Alternating left/right photo layout with gold diamond studs, sepia-toned archival photos, captions, and quote boxes
- Interactive bullet point cards with numbered maroon/gold badges
- Highlight banners on Slide 1 & Slide 10
- Smooth cross-fade slide transitions and staggered entrance animations
"""

import os
import io
import numpy as np
from PIL import Image, ImageEnhance
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml import parse_xml

# Color Palette Constants (Exact match to web app)
COLOR_LEATHER_BG = RGBColor(42, 26, 14)       # #2A1A0E - Dark Leather Canvas
COLOR_PARCHMENT = RGBColor(250, 244, 230)     # #FAF4E6 - Main Slide Parchment Fill
COLOR_CARD_PARCHMENT = RGBColor(241, 228, 195)# #F1E4C3 - Card Parchment Fill
COLOR_MAROON = RGBColor(122, 31, 31)          # #7A1F1F - Deep Maroon
COLOR_GOLD = RGBColor(184, 134, 43)           # #B8862B - Antique Gold
COLOR_BRIGHT_GOLD = RGBColor(212, 175, 55)    # #D4AF37 - Bright Accent Gold
COLOR_PALE_GOLD = RGBColor(212, 190, 136)     # #D4BE88 - Caption Gold Accent
COLOR_DARK_BROWN = RGBColor(61, 39, 23)       # #3D2717 - Dark Subtitle/Box Brown
COLOR_TEXT_DARK = RGBColor(42, 26, 14)        # #2A1A0E - Main Body Text Dark
COLOR_TEXT_LIGHT = RGBColor(250, 244, 230)    # #FAF4E6 - Light Text for dark boxes
COLOR_SUBTITLE_BROWN = RGBColor(74, 50, 32)   # #4A3220 - Subtitle & Chapter Brown

# Font Families with System Fallbacks
FONT_TITLE = "Cinzel Decorative"
FONT_HEADING = "Cinzel"
FONT_BODY = "EB Garamond"
FONT_CAPTION = "IM Fell English"
FONT_FALLBACK_SERIF = "Georgia"

def get_shape_id(shape):
    """Retrieve internal OpenXML shape ID for animation timing."""
    cNvPr = shape._element.find('.//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr')
    if cNvPr is not None:
        return cNvPr.get('id')
    return None

def process_image_sepia(image_path, target_width=1200):
    """
    Applies 35% sepia filter, 105% contrast, and 95% brightness
    matching the React application's exact CSS filter properties.
    """
    if not os.path.exists(image_path):
        return None

    img = Image.open(image_path).convert('RGB')

    # Resize if needed to maintain high quality while keeping size lean
    if img.width > target_width:
        ratio = target_width / img.width
        new_height = int(img.height * ratio)
        img = img.resize((target_width, new_height), Image.Resampling.LANCZOS)

    # Apply brightness & contrast
    enh_bri = ImageEnhance.Brightness(img)
    img = enh_bri.enhance(0.95)
    enh_con = ImageEnhance.Contrast(img)
    img = enh_con.enhance(1.05)

    # 35% Sepia Matrix transformation
    arr = np.array(img, dtype=np.float32)
    sepia_matrix = np.array([
        [0.393, 0.769, 0.189],
        [0.349, 0.686, 0.168],
        [0.272, 0.534, 0.131]
    ])
    sepia_arr = arr.dot(sepia_matrix.T)
    sepia_arr = np.clip(sepia_arr, 0, 255)

    blended = arr * 0.65 + sepia_arr * 0.35
    blended = np.clip(blended, 0, 255).astype(np.uint8)
    final_img = Image.fromarray(blended)

    # Save to BytesIO buffer
    buf = io.BytesIO()
    final_img.save(buf, format='JPEG', quality=88, optimize=True)
    buf.seek(0)
    return buf

# Slide Data Definition (Exact replica of src/components/gandhiData.ts)
SLIDES_DATA = [
  {
    "no": 1,
    "roman": 'I',
    "kicker": 'CHAPTER I • 21ST CENTURY INQUIRY',
    "title": 'IS GANDHIAN PHILOSOPHY STILL RELEVANT IN THE 21ST CENTURY?',
    "sub": 'An Exploration of Truth, Non-Violence, and Moral Courage in the Modern Age',
    "points": [
      'Gandhi Jayanti, 2 October — Commemorating the global heritage of peace and truth.',
      'Mohandas Karamchand Gandhi, 1869–1948 — Leader of India\'s non-violent freedom movement.',
      'Presented by Dev Vashisht — Examining the enduring relevance of Gandhian ideals today.'
    ],
    "photo": 'public/gandhi/gandhi-2.jpg',
    "caption": 'Mohandas Karamchand Gandhi (1869–1948) — Father of the Nation and apostle of non-violence.',
    "quote": 'My life is my message.',
    "highlight": 'Presented by Dev Vashisht'
  },
  {
    "no": 2,
    "roman": 'II',
    "kicker": 'CHAPTER II • IDENTITY & LEGACY',
    "title": 'WHO WAS GANDHI?',
    "sub": 'From a Young Barrister in London to the Leader of a Global Movement',
    "points": [
      'Born 2 October 1869 in Porbandar, Gujarat, India.',
      'Trained as a lawyer in London at the Inner Temple.',
      'Led India\'s historic non-violent freedom struggle against colonial rule.',
      'Called Mahatma (\'Great Soul\') and affectionately revered as Bapu (\'Father\').',
      'The central question: do his ideas and methods still work in today\'s complex world?'
    ],
    "photo": 'public/gandhi/gandhi-2.jpg',
    "caption": 'Gandhi, London, 1931',
    "quote": 'In a gentle way, you can shake the world.'
  },
  {
    "no": 3,
    "roman": 'III',
    "kicker": 'CHAPTER III • PHILOSOPHICAL PILLARS',
    "title": 'THE CORE IDEAS',
    "sub": 'The Timeless Foundations of Gandhian Philosophy and Ethics',
    "points": [
      'Satya (Truth) — Absolute adherence to truth and honesty in thought, word, and deed.',
      'Ahimsa (Non-violence) — Active love, compassion, and refraining from harm against any living being.',
      'Satyagraha (Peaceful resistance) — Fearless soul-force standing firm against injustice without malice.',
      'Sarvodaya (Welfare of all) — Universal upliftment prioritizing the most vulnerable and marginalized.',
      'Swadeshi (Self-reliance) — Fostering local production, community strength, and economic independence.',
      'Simple living — Voluntary simplicity, mindful consumption, and harmony with the natural world.'
    ],
    "photo": 'public/gandhi/gandhi-3.jpg',
    "caption": 'Gandhi and Kasturba, South Africa, 1902',
    "quote": 'Truth is the sovereign principle, which includes numerous other principles.'
  },
  {
    "no": 4,
    "roman": 'IV',
    "kicker": 'CHAPTER IV • 21ST CENTURY REALITIES',
    "title": 'THE WORLD WE LIVE IN',
    "sub": 'The Urgent Humanitarian, Social, and Ecological Crises of Our Era',
    "points": [
      'Wars and conflict: Escalating geopolitical hostilities, armed warfare, and global instability.',
      'Polarisation and hate online: Digital echo chambers, social fragmentation, and rising hostility.',
      'Climate crisis: Global environmental degradation, extreme weather, and resource depletion.',
      'Inequality: Widening socio-economic divides and unequal access to essential opportunities.',
      'Misinformation: Rapid algorithmic dissemination of falsehoods eroding public trust.',
      'Consumerism and waste: Hyper-materialistic lifestyles straining the planet\'s finite resources.'
    ],
    "photo": 'public/gandhi/gandhi-4.jpg',
    "caption": 'Gandhi and Nehru, Mumbai, 1946',
    "quote": 'The world will live in peace only when the individuals composing it make up their minds to do so.'
  },
  {
    "no": 5,
    "roman": 'V',
    "kicker": 'CHAPTER V • POWER OF NON-VIOLENCE',
    "title": 'AHIMSA IN ACTION',
    "sub": 'How Non-Violent Resistance Transformed Global Civil Rights and History',
    "points": [
      'The US civil rights movement (Martin Luther King Jr.) and South African anti-apartheid struggle (Nelson Mandela) drew deeply on Gandhian ideas.',
      'Chenoweth and Stephan\'s study (analyzing 323 campaigns from 1900 to 2006) found non-violent campaigns succeeded more often than violent ones.',
      '2 October is recognized internationally by the United Nations as the International Day of Non-Violence.'
    ],
    "photo": 'public/gandhi/gandhi-5.jpg',
    "caption": 'Gandhi with Abdul Ghaffar Khan, 1940',
    "quote": 'Non-violence is the greatest force at the disposal of mankind. — Mahatma Gandhi'
  },
  {
    "no": 6,
    "roman": 'VI',
    "kicker": 'CHAPTER VI • TRUTH IN THE DIGITAL AGE',
    "title": 'SATYA IN THE AGE OF MISINFORMATION',
    "sub": 'Practicing Discernment, Integrity, and Civil Dialogue in a Connected Society',
    "points": [
      'Fake news spreads fast: Digital algorithms amplify sensationalism and falsehoods at unprecedented speed.',
      'Verify before you speak or share: Upholding factual accuracy and critical thinking before disseminating information.',
      'Truth-telling as a daily discipline: Living with intellectual honesty and moral transparency.',
      'Dialogue instead of online outrage: Choosing constructive communication and empathy over reactionary anger.'
    ],
    "photo": 'public/gandhi/gandhi-6.jpg',
    "caption": 'The Salt March, 1930',
    "quote": 'Morality is the basis of things, and truth is the substance of all morality.'
  },
  {
    "no": 7,
    "roman": 'VII',
    "kicker": 'CHAPTER VII • ECOLOGICAL WISDOM',
    "title": 'SIMPLE LIVING AND THE PLANET',
    "sub": 'Sustainable Living, Mindful Consumption, and Climate Responsibility',
    "points": [
      '\'The earth provides enough for everyone\'s needs, but not everyone\'s greed\' (attributed to Gandhi).',
      'Direct links to modern sustainability, conscious minimalism, and circular resource use.',
      'India\'s Mission LiFE (Lifestyle for Environment) actively promotes individual and community eco-friendly lifestyles.'
    ],
    "photo": 'public/gandhi/gandhi-7.jpg',
    "caption": 'Gandhi at Dandi, 5 April 1930',
    "quote": 'The earth provides enough to satisfy every man\'s needs, but not every man\'s greed.'
  },
  {
    "no": 8,
    "roman": 'VIII',
    "kicker": 'CHAPTER VIII • LOCAL ECONOMIES & DIGNITY',
    "title": 'SWADESHI AND SARVODAYA TODAY',
    "sub": 'Grassroots Empowerment, Inclusive Growth, and Universal Cleanliness',
    "points": [
      'Local economies and self-reliance: Strengthening local supply networks and supporting homegrown enterprise.',
      'Khadi and village industries: Empowering rural artisans and promoting eco-conscious handloom textiles.',
      'Inclusion and dignity for all: Ensuring that progress uplift the most disadvantaged in society.',
      'Swachh Bharat Abhiyan: A nationwide sanitation and cleanliness movement launched on Gandhi Jayanti (2 Oct 2014).'
    ],
    "photo": 'public/gandhi/gandhi-8.jpg',
    "caption": 'Gandhi and Nehru, Quit India session, 1942',
    "quote": 'Recall the face of the poorest and the weakest person you have seen, and ask if your step will be of any use to them.'
  },
  {
    "no": 9,
    "roman": 'IX',
    "kicker": 'CHAPTER IX • CRITICAL PERSPECTIVE',
    "title": 'LIMITS AND CRITICISMS',
    "sub": 'Nuance, Context, and Thoughtful Adaptation in the Modern Era',
    "points": [
      'Non-violence alone may not be enough when confronting extreme violence and ruthless authoritarian regimes.',
      'Some of Gandhi\'s views (on caste and race during his early South Africa years) are legitimately criticised.',
      'Ideas need thoughtful application and dynamic adaptation, not rigid or blind copying.'
    ],
    "photo": 'public/gandhi/gandhi-9.jpg',
    "caption": 'Portrait, late 1930s',
    "quote": 'I want the cultures of all lands to be blown about my house as freely as possible, but I refuse to be blown off my feet.'
  },
  {
    "no": 10,
    "roman": 'X',
    "kicker": 'CHAPTER X • LIVING RELEVANCE',
    "title": 'THE VERDICT: STILL RELEVANT',
    "sub": 'A Living Ethical Compass for Contemporary Life and Global Citizenship',
    "points": [
      'A compass, not a rulebook: An enduring framework for moral decision-making in personal and public affairs.',
      'Choose peace, truth, and simplicity in everyday choices, leadership, and community action.',
      '"Be the change you wish to see in the world" (attributed to Gandhi) — Individual integrity inspires collective transformation.',
      'Thank you! Jai Hind!'
    ],
    "photo": 'public/gandhi/gandhi-10.jpg',
    "caption": 'Gandhi, London, 1931',
    "quote": 'You must be the change you wish to see in the world.',
    "highlight": 'Thank you! Jai Hind!',
    "credit": 'Photos: Wikimedia Commons (public domain)'
  }
]

def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    total_slides = len(SLIDES_DATA)
    total_roman = "X"

    for idx, slide_data in enumerate(SLIDES_DATA):
        slide = prs.slides.add_slide(blank_layout)
        animated_shape_ids = []

        is_photo_left = (slide_data["no"] % 2 == 0)

        # ----------------------------------------------------
        # 1. Slide Canvas Background: Dark Leather Brown (#2A1A0E)
        # ----------------------------------------------------
        bg_shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            0, 0, Inches(13.333), Inches(7.5)
        )
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = COLOR_LEATHER_BG
        bg_shape.line.color.rgb = COLOR_LEATHER_BG

        # ----------------------------------------------------
        # 2. Main Parchment Sheet Container (#FAF4E6)
        # ----------------------------------------------------
        sheet_left = Inches(0.4)
        sheet_top = Inches(0.28)
        sheet_width = Inches(12.533)
        sheet_height = Inches(6.94)

        sheet_shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            sheet_left, sheet_top, sheet_width, sheet_height
        )
        sheet_shape.fill.solid()
        sheet_shape.fill.fore_color.rgb = COLOR_PARCHMENT
        # Deep Maroon Outer Border (3pt)
        sheet_shape.line.color.rgb = COLOR_MAROON
        sheet_shape.line.width = Pt(3)

        # Inner Antique Gold Decorative Border Line
        inner_inset = Inches(0.07)
        inner_border = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            sheet_left + inner_inset,
            sheet_top + inner_inset,
            sheet_width - (inner_inset * 2),
            sheet_height - (inner_inset * 2)
        )
        inner_border.fill.background()
        inner_border.line.color.rgb = COLOR_GOLD
        inner_border.line.width = Pt(0.75)

        # 4 Ornate Corner Diamonds (❖)
        corner_coords = [
            (sheet_left + Inches(0.12), sheet_top + Inches(0.1)),
            (sheet_left + sheet_width - Inches(0.32), sheet_top + Inches(0.1)),
            (sheet_left + Inches(0.12), sheet_top + sheet_height - Inches(0.3)),
            (sheet_left + sheet_width - Inches(0.32), sheet_top + sheet_height - Inches(0.3)),
        ]
        for cx, cy in corner_coords:
            c_box = slide.shapes.add_textbox(cx, cy, Inches(0.25), Inches(0.25))
            tf = c_box.text_frame
            tf.word_wrap = False
            tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
            p = tf.paragraphs[0]
            p.text = "❖"
            p.font.name = FONT_HEADING
            p.font.size = Pt(11)
            p.font.color.rgb = COLOR_GOLD

        # ----------------------------------------------------
        # 3. Header Area: Chapter Kicker + Folio Badge + Title + Subtitle
        # ----------------------------------------------------
        header_left = sheet_left + Inches(0.3)
        header_width = sheet_width - Inches(0.6)
        header_top = sheet_top + Inches(0.2)

        # 3a. Top Header Bar: Kicker on Left, Folio Badge on Right
        kicker_box = slide.shapes.add_textbox(header_left, header_top, Inches(8.5), Inches(0.32))
        ktf = kicker_box.text_frame
        ktf.word_wrap = False
        ktf.margin_left = ktf.margin_top = ktf.margin_right = ktf.margin_bottom = 0
        kp = ktf.paragraphs[0]
        kp.text = f"❖  {slide_data['kicker']}"
        kp.font.name = FONT_HEADING
        kp.font.size = Pt(10.5)
        kp.font.bold = True
        kp.font.color.rgb = COLOR_MAROON

        # Folio Badge
        badge_width = Inches(2.2)
        badge_height = Inches(0.3)
        badge_left = header_left + header_width - badge_width
        badge_shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            badge_left, header_top - Inches(0.02), badge_width, badge_height
        )
        badge_shape.fill.solid()
        badge_shape.fill.fore_color.rgb = COLOR_CARD_PARCHMENT
        badge_shape.line.color.rgb = COLOR_GOLD
        badge_shape.line.width = Pt(1)
        btf = badge_shape.text_frame
        btf.margin_left = btf.margin_top = btf.margin_right = btf.margin_bottom = 0
        bp = btf.paragraphs[0]
        bp.alignment = PP_ALIGN.CENTER
        bp.text = f"FOLIO {slide_data['roman']} DE {total_roman}"
        bp.font.name = FONT_HEADING
        bp.font.size = Pt(9.5)
        bp.font.bold = True
        bp.font.color.rgb = COLOR_SUBTITLE_BROWN

        # 3b. Title
        title_top = header_top + Inches(0.34)
        title_box = slide.shapes.add_textbox(header_left, title_top, header_width, Inches(0.65))
        ttf = title_box.text_frame
        ttf.word_wrap = True
        ttf.margin_left = ttf.margin_top = ttf.margin_right = ttf.margin_bottom = 0
        tp = ttf.paragraphs[0]
        tp.text = slide_data['title']
        tp.font.name = FONT_TITLE
        # Dynamic font size based on length
        title_len = len(slide_data['title'])
        if title_len > 45:
            tp.font.size = Pt(20)
        elif title_len > 30:
            tp.font.size = Pt(23)
        else:
            tp.font.size = Pt(26)
        tp.font.bold = True
        tp.font.color.rgb = COLOR_TEXT_DARK

        # 3c. Thin Gold Divider Line with ❦
        div_top = title_top + Inches(0.68)
        div_line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            header_left, div_top + Inches(0.08), header_width, Inches(0.015)
        )
        div_line.fill.solid()
        div_line.fill.fore_color.rgb = COLOR_GOLD
        div_line.line.color.rgb = COLOR_GOLD

        div_center = slide.shapes.add_textbox(header_left + (header_width / 2) - Inches(0.3), div_top - Inches(0.02), Inches(0.6), Inches(0.22))
        dtf = div_center.text_frame
        dtf.margin_left = dtf.margin_top = dtf.margin_right = dtf.margin_bottom = 0
        dp = dtf.paragraphs[0]
        dp.alignment = PP_ALIGN.CENTER
        dp.text = "❦"
        dp.font.name = FONT_FALLBACK_SERIF
        dp.font.size = Pt(11)
        dp.font.color.rgb = COLOR_MAROON

        # 3d. Subtitle
        sub_top = div_top + Inches(0.18)
        sub_box = slide.shapes.add_textbox(header_left, sub_top, header_width, Inches(0.35))
        stf = sub_box.text_frame
        stf.word_wrap = True
        stf.margin_left = stf.margin_top = stf.margin_right = stf.margin_bottom = 0
        sp = stf.paragraphs[0]
        sp.text = slide_data['sub']
        sp.font.name = FONT_BODY
        sp.font.size = Pt(13)
        sp.font.italic = True
        sp.font.color.rgb = COLOR_DARK_BROWN

        # ----------------------------------------------------
        # 4. Main Body: Alternating Content vs Photo Layout
        # ----------------------------------------------------
        body_top = sub_top + Inches(0.42)
        body_height = Inches(4.32)
        gap = Inches(0.35)

        content_width = Inches(7.1)
        photo_width = header_width - content_width - gap

        if is_photo_left:
            photo_col_left = header_left
            content_col_left = header_left + photo_width + gap
        else:
            content_col_left = header_left
            photo_col_left = header_left + content_width + gap

        # ====================================================
        # 4A. PHOTO & QUOTE COLUMN
        # ====================================================
        # Photo Outer Container Frame (#2A1A0E background, #B8862B 1.5pt border)
        photo_frame = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            photo_col_left, body_top, photo_width, body_height
        )
        photo_frame.fill.solid()
        photo_frame.fill.fore_color.rgb = COLOR_LEATHER_BG
        photo_frame.line.color.rgb = COLOR_GOLD
        photo_frame.line.width = Pt(1.5)

        # Gold Diamond Corner Studs on Photo Frame
        stud_size = Inches(0.12)
        stud_coords = [
            (photo_col_left + Inches(0.04), body_top + Inches(0.04)),
            (photo_col_left + photo_width - Inches(0.16), body_top + Inches(0.04)),
            (photo_col_left + Inches(0.04), body_top + body_height - Inches(0.16)),
            (photo_col_left + photo_width - Inches(0.16), body_top + body_height - Inches(0.16))
        ]
        for sx, sy in stud_coords:
            stud = slide.shapes.add_shape(MSO_SHAPE.DIAMOND, sx, sy, stud_size, stud_size)
            stud.fill.solid()
            stud.fill.fore_color.rgb = COLOR_BRIGHT_GOLD
            stud.line.color.rgb = COLOR_MAROON
            stud.line.width = Pt(0.5)

        # Embed Sepia-Toned Historical Photo
        img_inset = Inches(0.14)
        img_left = photo_col_left + img_inset
        img_top = body_top + img_inset
        img_w = photo_width - (img_inset * 2)

        # Calculate image height based on whether quote box is present
        has_quote = bool(slide_data.get('quote'))
        if has_quote:
            img_h = Inches(2.25)
        else:
            img_h = Inches(2.95)

        photo_path = slide_data['photo']
        processed_img_buf = process_image_sepia(photo_path)
        if processed_img_buf:
            slide.shapes.add_picture(processed_img_buf, img_left, img_top, img_w, img_h)

        # Archive Header Ribbon on top-left of image
        ribbon_w = Inches(1.7)
        ribbon_h = Inches(0.24)
        ribbon = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            img_left + Inches(0.08), img_top + Inches(0.08), ribbon_w, ribbon_h
        )
        ribbon.fill.solid()
        ribbon.fill.fore_color.rgb = COLOR_LEATHER_BG
        ribbon.line.color.rgb = COLOR_GOLD
        ribbon.line.width = Pt(0.75)
        rtf = ribbon.text_frame
        rtf.margin_left = rtf.margin_top = rtf.margin_right = rtf.margin_bottom = 0
        rp = rtf.paragraphs[0]
        rp.alignment = PP_ALIGN.CENTER
        rp.text = f"ARCHIVE • CH. {slide_data['roman']}"
        rp.font.name = FONT_HEADING
        rp.font.size = Pt(7.5)
        rp.font.bold = True
        rp.font.color.rgb = COLOR_CARD_PARCHMENT

        # Photo Caption
        caption_top = img_top + img_h + Inches(0.08)
        caption_box = slide.shapes.add_textbox(img_left, caption_top, img_w, Inches(0.5))
        ctf = caption_box.text_frame
        ctf.word_wrap = True
        ctf.margin_left = ctf.margin_top = ctf.margin_right = ctf.margin_bottom = 0
        cp = ctf.paragraphs[0]
        cp.text = slide_data['caption']
        cp.font.name = FONT_CAPTION
        cp.font.size = Pt(10)
        cp.font.italic = True
        cp.font.color.rgb = COLOR_TEXT_LIGHT

        if slide_data.get('credit'):
            cp2 = ctf.add_paragraph()
            cp2.text = f"❦ {slide_data['credit']}"
            cp2.font.name = FONT_CAPTION
            cp2.font.size = Pt(8.5)
            cp2.font.italic = True
            cp2.font.color.rgb = COLOR_PALE_GOLD

        # Dictum Quote Card
        if has_quote:
            quote_top = body_top + body_height - Inches(1.22)
            quote_h = Inches(1.1)
            quote_box = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                img_left, quote_top, img_w, quote_h
            )
            quote_box.fill.solid()
            quote_box.fill.fore_color.rgb = COLOR_DARK_BROWN
            quote_box.line.color.rgb = COLOR_GOLD
            quote_box.line.width = Pt(1)

            qtf = quote_box.text_frame
            qtf.word_wrap = True
            qtf.margin_left = Inches(0.12)
            qtf.margin_right = Inches(0.12)
            qtf.margin_top = Inches(0.08)
            qtf.margin_bottom = Inches(0.06)

            qp_head = qtf.paragraphs[0]
            qp_head.text = "❝ DICTUM MAHATMA GANDHI"
            qp_head.font.name = FONT_HEADING
            qp_head.font.size = Pt(8.5)
            qp_head.font.bold = True
            qp_head.font.color.rgb = COLOR_BRIGHT_GOLD

            qp_text = qtf.add_paragraph()
            qp_text.text = f'"{slide_data["quote"]}"'
            qp_text.font.name = FONT_CAPTION
            qp_text.font.size = Pt(10.5)
            qp_text.font.italic = True
            qp_text.font.color.rgb = COLOR_CARD_PARCHMENT

        # ====================================================
        # 4B. CONTENT / BULLET CARDS COLUMN
        # ====================================================
        points = slide_data['points']
        num_points = len(points)
        has_highlight = bool(slide_data.get('highlight'))

        # Calculate heights for bullet cards
        available_content_h = body_height
        if has_highlight:
            highlight_h = Inches(0.55)
            available_bullet_h = available_content_h - highlight_h - Inches(0.15)
        else:
            available_bullet_h = available_content_h

        card_gap = Inches(0.1)
        card_h = (available_bullet_h - (card_gap * (num_points - 1))) / num_points
        card_h = min(card_h, Inches(1.1))  # Cap maximum height

        for p_idx, point_text in enumerate(points):
            cur_card_top = body_top + (p_idx * (card_h + card_gap))

            # Card Outer Box (#F1E4C3 with #B8862B border)
            card_shape = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                content_col_left, cur_card_top, content_width, card_h
            )
            card_shape.fill.solid()
            card_shape.fill.fore_color.rgb = COLOR_CARD_PARCHMENT
            card_shape.line.color.rgb = COLOR_GOLD
            card_shape.line.width = Pt(0.75)

            # Record shape ID for sequential fade-in animation
            sp_id = get_shape_id(card_shape)
            if sp_id:
                animated_shape_ids.append(sp_id)

            # Number Badge inside Card
            badge_sz = min(Inches(0.38), card_h - Inches(0.16))
            badge_top_inset = cur_card_top + Inches(0.08)
            num_badge = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                content_col_left + Inches(0.12), badge_top_inset, badge_sz, badge_sz
            )
            num_badge.fill.solid()
            num_badge.fill.fore_color.rgb = COLOR_MAROON
            num_badge.line.color.rgb = COLOR_GOLD
            num_badge.line.width = Pt(1)
            ntf = num_badge.text_frame
            ntf.margin_left = ntf.margin_top = ntf.margin_right = ntf.margin_bottom = 0
            np_text = ntf.paragraphs[0]
            np_text.alignment = PP_ALIGN.CENTER
            np_text.text = str(p_idx + 1)
            np_text.font.name = FONT_HEADING
            np_text.font.size = Pt(11)
            np_text.font.bold = True
            np_text.font.color.rgb = COLOR_CARD_PARCHMENT

            # Bullet Text
            text_left = content_col_left + Inches(0.12) + badge_sz + Inches(0.12)
            text_w = content_width - (text_left - content_col_left) - Inches(0.12)

            tb = slide.shapes.add_textbox(text_left, cur_card_top + Inches(0.06), text_w, card_h - Inches(0.12))
            ptf = tb.text_frame
            ptf.word_wrap = True
            ptf.margin_left = ptf.margin_top = ptf.margin_right = ptf.margin_bottom = 0
            ptf_p = ptf.paragraphs[0]
            ptf_p.text = point_text
            ptf_p.font.name = FONT_BODY

            # Adjust font size dynamically to fit
            if num_points >= 6:
                ptf_p.font.size = Pt(11)
            elif num_points >= 5:
                ptf_p.font.size = Pt(12)
            else:
                ptf_p.font.size = Pt(13)

            ptf_p.font.color.rgb = COLOR_TEXT_DARK

        # Highlight Banner (Slide 1 & Slide 10)
        if has_highlight:
            hl_top = body_top + body_height - Inches(0.55)
            hl_shape = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                content_col_left, hl_top, content_width, Inches(0.55)
            )
            hl_shape.fill.solid()
            hl_shape.fill.fore_color.rgb = COLOR_LEATHER_BG
            hl_shape.line.color.rgb = COLOR_GOLD
            hl_shape.line.width = Pt(1.5)

            hl_tf = hl_shape.text_frame
            hl_tf.word_wrap = True
            hl_tf.margin_left = Inches(0.15)
            hl_tf.margin_right = Inches(0.15)
            hl_tf.margin_top = Inches(0.08)
            hl_tf.margin_bottom = Inches(0.08)

            hl_p = hl_tf.paragraphs[0]
            hl_p.text = f"❖  {slide_data['highlight'].upper()}       —       IN MEMORIAM MAHATMA GANDHI"
            hl_p.font.name = FONT_HEADING
            hl_p.font.size = Pt(10.5)
            hl_p.font.bold = True
            hl_p.font.color.rgb = COLOR_BRIGHT_GOLD

        # ----------------------------------------------------
        # 5. Ornate Footer Row
        # ----------------------------------------------------
        footer_top = sheet_top + sheet_height - Inches(0.48)
        footer_w = header_width

        # Thin Maroon Separator Line
        ft_line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            header_left, footer_top, footer_w, Inches(0.015)
        )
        ft_line.fill.solid()
        ft_line.fill.fore_color.rgb = COLOR_MAROON
        ft_line.line.color.rgb = COLOR_MAROON

        # Left Footer
        f_left_box = slide.shapes.add_textbox(header_left, footer_top + Inches(0.06), Inches(3.8), Inches(0.26))
        fl_tf = f_left_box.text_frame
        fl_tf.margin_left = fl_tf.margin_top = fl_tf.margin_right = fl_tf.margin_bottom = 0
        fl_p = fl_tf.paragraphs[0]
        fl_p.text = "❖  GANDHI JAYANTI • II OCTOBER"
        fl_p.font.name = FONT_HEADING
        fl_p.font.size = Pt(8.5)
        fl_p.font.bold = True
        fl_p.font.color.rgb = COLOR_MAROON

        # Center Footer
        f_center_box = slide.shapes.add_textbox(header_left + Inches(3.8), footer_top + Inches(0.06), Inches(4.5), Inches(0.26))
        fc_tf = f_center_box.text_frame
        fc_tf.margin_left = fc_tf.margin_top = fc_tf.margin_right = fc_tf.margin_bottom = 0
        fc_p = fc_tf.paragraphs[0]
        fc_p.alignment = PP_ALIGN.CENTER
        fc_p.text = "TRUTH  •  NON-VIOLENCE  •  PEACE  •  DIGNITY"
        fc_p.font.name = FONT_HEADING
        fc_p.font.size = Pt(8.5)
        fc_p.font.bold = True
        fc_p.font.color.rgb = COLOR_MAROON

        # Right Footer
        f_right_box = slide.shapes.add_textbox(header_left + footer_w - Inches(2.2), footer_top + Inches(0.06), Inches(2.2), Inches(0.26))
        fr_tf = f_right_box.text_frame
        fr_tf.margin_left = fr_tf.margin_top = fr_tf.margin_right = fr_tf.margin_bottom = 0
        fr_p = fr_tf.paragraphs[0]
        fr_p.alignment = PP_ALIGN.RIGHT
        fr_p.text = f"FOLIO {slide_data['roman']} / {total_roman}  ❖"
        fr_p.font.name = FONT_HEADING
        fr_p.font.size = Pt(8.5)
        fr_p.font.bold = True
        fr_p.font.color.rgb = COLOR_SUBTITLE_BROWN

        # ----------------------------------------------------
        # 6. Smooth Slide Transitions & Sequential Animations (OpenXML)
        # ----------------------------------------------------
        # Slide Cross-Fade Transition
        transition_xml = parse_xml(
            '<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" spd="med">'
            '<p:fade/>'
            '</p:transition>'
        )
        slide._element.append(transition_xml)

        # Sequential Fade-in Animations for Bullet Points
        if animated_shape_ids:
            child_nodes = []
            curr_id = 3
            for spid in animated_shape_ids:
                c1 = curr_id
                c2 = curr_id + 1
                curr_id += 2
                child_nodes.append(f'''
                <p:par>
                  <p:cTn id="{c1}" fill="hold">
                    <p:stCondLst>
                      <p:cond delay="0"/>
                    </p:stCondLst>
                    <p:childTnLst>
                      <p:animEffect transition="in" filter="fade">
                        <p:cBhvr>
                          <p:cTn id="{c2}" dur="350"/>
                          <p:tgtEl>
                            <p:spTgt spid="{spid}"/>
                          </p:tgtEl>
                        </p:cBhvr>
                      </p:animEffect>
                    </p:childTnLst>
                  </p:cTn>
                </p:par>
                ''')

            timing_str = f'''
            <p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
              <p:tnLst>
                <p:par>
                  <p:cTn id="1" dur="indefinite" restart="always" nodeType="tmRoot">
                    <p:childTnLst>
                      <p:seq concurrent="1" nextAc="seek">
                        <p:cTn id="2" dur="indefinite" nodeType="mainSeq">
                          <p:childTnLst>
                            {''.join(child_nodes)}
                          </p:childTnLst>
                        </p:cTn>
                        <p:prevCondLst>
                          <p:cond evt="onPrev" delay="0">
                            <p:tgtEl>
                              <p:spTgt spid="0"/>
                            </p:tgtEl>
                          </p:cond>
                        </p:prevCondLst>
                        <p:nextCondLst>
                          <p:cond evt="onNext" delay="0">
                            <p:tgtEl>
                              <p:spTgt spid="0"/>
                            </p:tgtEl>
                          </p:cond>
                        </p:nextCondLst>
                      </p:seq>
                    </p:childTnLst>
                  </p:cTn>
                </p:par>
              </p:tnLst>
            </p:timing>
            '''
            timing_xml = parse_xml(timing_str)
            slide._element.append(timing_xml)

    output_path = "gandhi-presentation.pptx"
    prs.save(output_path)
    print(f"Presentation generated successfully at: {output_path} ({os.path.getsize(output_path)/1024:.1f} KB)")

    # Also copy to public and dist folders
    for dest_dir in ["public", "dist"]:
        if os.path.exists(dest_dir):
            dest_file = os.path.join(dest_dir, "gandhi-presentation.pptx")
            with open(output_path, "rb") as src, open(dest_file, "wb") as dst:
                dst.write(src.read())
            print(f"Copied to: {dest_file}")

if __name__ == "__main__":
    build_presentation()
